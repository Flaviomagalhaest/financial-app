"""Exports fictícios do Mobills, no formato do `## Answer` dos tickets 01 e 07: o geral com a aba
`Receitas e Despesas` (8 colunas, data em texto, valor com sinal, linha em branco e linha de total
no fim) e as faturas por cartão e mês (`<cartão>_-_<AAAAMM>.xlsx`, aba `Receitas`, 4 colunas, valor
em texto `R$ 1.234,56`). Dado real nunca entra aqui (ADR 0007)."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUNAS_DO_GERAL = [
    "Data",
    "Descrição",
    "Valor",
    "Conta",
    "Situação",
    "Categoria",
    "Subcategoria",
    "Tags",
]


def linha(data, descricao, valor, categoria, situacao="Paga", tags=""):
    """Uma linha do geral: `valor` em reais com sinal (despesa negativa)."""
    return [data, descricao, valor, "Carteira", situacao, categoria, "", tags]


def escrever_geral(pasta: Path, nome: str, linhas: list[list]) -> Path:
    pasta.mkdir(parents=True, exist_ok=True)
    livro = Workbook()
    aba = livro.active
    aba.title = "Receitas e Despesas"
    aba.append(COLUNAS_DO_GERAL)
    for valores in linhas:
        aba.append(valores)
    aba.append([])
    aba.append(["Total (receitas - despesas)", None, round(sum(v[2] for v in linhas), 2)])
    for outra in ("Despesas", "Receitas", "Transferências"):
        livro.create_sheet(outra).append(["Data", "Descrição", "Valor"])
    caminho = pasta / nome
    livro.save(caminho)
    return caminho


def reais(valor: float) -> str:
    """O texto do valor na fatura: `R$ 1.234,56`, com `-` antes nos créditos."""
    texto = f"{abs(valor):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{'-' if valor < 0 else ''}R$ {texto}"


def escrever_fatura(pasta: Path, cartao: str, rotulo: str, linhas: list[tuple]) -> Path:
    """Linhas `(data, descrição, categoria, valor)`, valor positivo na compra."""
    destino = pasta / cartao
    destino.mkdir(parents=True, exist_ok=True)
    livro = Workbook()
    aba = livro.active
    aba.title = "Receitas"
    aba.append(["Data", "Descrição", "Categoria", "Valor"])
    for data, descricao, categoria, valor in linhas:
        aba.append([data, descricao, categoria, reais(valor)])
    caminho = destino / f"{cartao}_-_{rotulo}.xlsx"
    livro.save(caminho)
    return caminho


def ler_aba(planilha: Path, nome: str) -> list[dict]:
    """As linhas de uma aba da planilha de revisão, pelo cabeçalho."""
    aba = load_workbook(planilha)[nome]
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas)
    return [dict(zip(cabecalho, valores)) for valores in linhas]


def editar_aba(planilha: Path, nome: str, chave: str, valor_da_chave, coluna: str, novo) -> None:
    """Muda uma célula como o dono faria: na linha em que `chave` vale `valor_da_chave`."""
    livro = load_workbook(planilha)
    aba = livro[nome]
    cabecalho = [c.value for c in aba[1]]
    for celulas in aba.iter_rows(min_row=2):
        if celulas[cabecalho.index(chave)].value == valor_da_chave:
            celulas[cabecalho.index(coluna)].value = novo
            break
    else:
        raise AssertionError(f"{nome}: sem linha com {chave} = {valor_da_chave!r}")
    livro.save(planilha)
