"""A planilha de revisão (`## Answer` do ticket 18): o que o importador não decide sozinho, para o
dono editar e rodar de novo, e as abas de conferência contra o Mobills."""

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.mobills import de_para
from app.mobills.conferencia import Fatura, Mes
from app.mobills.cruzamento import Cruzamento
from app.mobills.leitura import centavos, data, texto
from app.mobills.parcelas import Parcelamento

FORMATO_DE_REAIS = "#,##0.00"


def reais(centavos: int) -> float:
    return centavos / 100


def reais_br(centavos: int) -> str:
    """`-1.234,56`, para o texto do "Por quê"."""
    texto = f"{abs(centavos) / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-{texto}" if centavos < 0 else texto


def data_br(dia: date) -> str:
    return f"{dia:%d/%m/%Y}"


COLUNA_DA_PASTA = "Pasta do cartão"
COLUNA_DA_DECISAO = "Decisão (juntar ou separar)"
COLUNA_DO_PAPEL = "Papel (categoria, tag, pessoa ou descartar)"
COLUNA_DO_CARTAO_ESCOLHIDO = "Cartão (escolha)"
COLUNA_DO_TIPO = "Tipo (corrente, poupanca ou carteira)"
COLUNA_DO_ARQUIVADO = "Arquivado (sim ou não)"
DECISOES_DO_PROVAVEL = {"juntar", "separar"}


@dataclass
class Decisoes:
    """O que o dono editou na planilha da rodada anterior."""

    provaveis: dict[tuple, str] = field(default_factory=dict)
    conta: tuple[str, str] | None = None  # nome e tipo
    # Por pasta, os campos de `Cartao`; o dia vem como o dono digitou, para a conferência dizer.
    cartoes: dict[str, dict] = field(default_factory=dict)
    de_para: dict[de_para.Chave, tuple[str, str]] = field(default_factory=dict)  # papel e nome
    cartao_dos_parcelamentos: dict[tuple, str] = field(default_factory=dict)
    cartao_dos_ajustes: dict[tuple[str, int], str] = field(default_factory=dict)  # mês e valor
    mobills_por_mes: dict[str, int] = field(default_factory=dict)
    mobills_por_fatura: dict[tuple[str, str], int] = field(default_factory=dict)


def ler_decisoes(caminho: Path) -> Decisoes:
    decisoes = Decisoes()
    if not caminho.exists():
        return decisoes
    livro = load_workbook(caminho, data_only=True)

    def valor(linha: dict, coluna: str) -> int | None:
        conteudo = linha.get(coluna)
        if texto(conteudo) == "":
            return None
        return centavos(conteudo, f"{caminho.name}, coluna {coluna}")

    for linha in _linhas(livro, "Conta")[:1]:
        decisoes.conta = (texto(linha.get("Nome")), texto(linha.get(COLUNA_DO_TIPO)).casefold())
    for linha in _linhas(livro, "Cartões"):
        decisoes.cartoes[texto(linha.get("Pasta"))] = {
            "nome": texto(linha.get("Cartão")),
            "arquivado": texto(linha.get(COLUNA_DO_ARQUIVADO)).casefold() == "sim",
            "fechamento": _dia_do_mes(linha.get("Dia do fechamento")),
            "vencimento": _dia_do_mes(linha.get("Dia do vencimento")),
        }
    for linha in _linhas(livro, "De-para"):
        chave = (
            texto(linha.get("Tipo")),
            texto(linha.get("No Mobills")),
            texto(linha.get("Nome no Mobills")),
        )
        decisoes.de_para[chave] = (
            texto(linha.get(COLUNA_DO_PAPEL)).casefold(),
            texto(linha.get("Nome no app")),
        )
    for linha in _linhas(livro, "Parcelamentos sem cartão"):
        chave = (
            texto(linha.get("Descrição")),
            valor(linha, "Valor"),
            texto(linha.get("Parcelas")).split(" ")[0],
            texto(linha.get("Primeiro mês")),
        )
        decisoes.cartao_dos_parcelamentos[chave] = texto(linha.get(COLUNA_DO_CARTAO_ESCOLHIDO))
    for linha in _linhas(livro, "Ajustes sem cartão"):
        chave = (texto(linha.get("Mês")), valor(linha, "Valor"))
        decisoes.cartao_dos_ajustes[chave] = texto(linha.get(COLUNA_DO_CARTAO_ESCOLHIDO))
    for linha in _linhas(livro, "Conferência"):
        if (do_mobills := valor(linha, "Mobills")) is not None:
            decisoes.mobills_por_mes[texto(linha.get("Mês"))] = do_mobills
    for linha in _linhas(livro, "Faturas"):
        if (do_mobills := valor(linha, "Mobills")) is not None:
            chave = (texto(linha.get("Cartão")), texto(linha.get("Mês")))
            decisoes.mobills_por_fatura[chave] = do_mobills
    for linha in _linhas(livro, "Pares prováveis"):
        decisao = texto(linha.get(COLUNA_DA_DECISAO)).casefold()
        if decisao in DECISOES_DO_PROVAVEL:
            chave = (
                texto(linha.get(COLUNA_DA_PASTA)),
                texto(linha.get("Mês")),
                data(linha.get("Data da compra")),
                texto(linha.get("Descrição na fatura")),
                valor(linha, "Valor"),
                data(linha.get("Data no geral")),
                texto(linha.get("Descrição no geral")),
            )
            decisoes.provaveis[chave] = decisao
    return decisoes


def _linhas(livro, nome: str) -> list[dict]:
    if nome not in livro.sheetnames:
        return []
    linhas = livro[nome].iter_rows(values_only=True)
    cabecalho = [texto(c) for c in next(linhas, ())]
    return [
        dict(zip(cabecalho, valores)) for valores in linhas if any(v is not None for v in valores)
    ]


def _dia_do_mes(valor) -> int | str | None:
    """O dia em número; o que não for número volta como o dono digitou."""
    if texto(valor) == "":
        return None
    try:
        return int(float(valor))
    except ValueError:
        return texto(valor)


def aba_conta(conta) -> tuple[list[str], list[list]]:
    return ["Nome", COLUNA_DO_TIPO], [[conta.nome, conta.tipo]]


def aba_cartoes(cartoes, pastas: dict[str, list[str]]) -> tuple[list[str], list[list]]:
    cabecalho = [
        "Pasta",
        "Cartão",
        "Primeira fatura",
        "Última fatura",
        COLUNA_DO_ARQUIVADO,
        "Dia do fechamento",
        "Dia do vencimento",
    ]
    linhas = [
        [
            c.pasta,
            c.nome,
            pastas[c.pasta][0] if pastas[c.pasta] else None,
            pastas[c.pasta][-1] if pastas[c.pasta] else None,
            "sim" if c.arquivado else "não",
            c.fechamento,
            c.vencimento,
        ]
        for c in cartoes
    ]
    return cabecalho, linhas


def aba_de_para(
    contagem: dict[de_para.Chave, int], regras: dict[de_para.Chave, de_para.Regra]
) -> tuple[list[str], list[list]]:
    cabecalho = [
        "Tipo",
        "No Mobills",
        "Nome no Mobills",
        "Lançamentos",
        COLUNA_DO_PAPEL,
        "Nome no app",
    ]
    linhas = [
        [*chave, contagem[chave], regras[chave].papel, regras[chave].nome]
        for chave in sorted(contagem)
    ]
    return cabecalho, linhas


def aba_pares_provaveis(cruzamento: Cruzamento) -> tuple[list[str], list[list]]:
    cabecalho = [
        COLUNA_DA_PASTA,
        "Mês",
        "Data da compra",
        "Descrição na fatura",
        "Categoria na fatura",
        "Valor",
        "Data no geral",
        "Descrição no geral",
        "Categoria no geral",
        "Tags no geral",
        COLUNA_DA_DECISAO,
    ]
    linhas = [
        [
            p.fatura.pasta,
            p.fatura.mes,
            data_br(p.fatura.data),
            p.fatura.descricao,
            p.fatura.categoria,
            reais(p.fatura.valor_centavos),
            data_br(p.geral.data),
            p.geral.descricao,
            p.geral.categoria,
            ", ".join(p.geral.tags),
            p.decisao,
        ]
        for p in cruzamento.provaveis
    ]
    return cabecalho, linhas


def aba_grupos_ambiguos(cruzamento: Cruzamento) -> tuple[list[str], list[list]]:
    """Um par por linha; o lado que sobrou fica em branco."""
    cabecalho = [
        COLUNA_DA_PASTA,
        "Mês",
        "Descrição",
        "Valor",
        "Categoria",
        "Data da compra",
        "Arquivo do geral",
        "Linha no arquivo",
        "Tags no geral",
        "Situação no geral",
    ]
    linhas = []
    for grupo in cruzamento.ambiguos:
        geral = next(g for _, g in grupo if g is not None)
        for f, g in grupo:
            linhas.append(
                [
                    f.pasta if f else None,
                    geral.mes,
                    geral.descricao,
                    reais(-geral.valor_centavos),
                    geral.categoria,
                    data_br(f.data) if f else None,
                    g.arquivo if g else None,
                    g.linha if g else None,
                    ", ".join(g.tags) if g else None,
                    g.situacao if g else None,
                ]
            )
    return cabecalho, linhas


def chave_do_parcelamento(parcelamento: Parcelamento) -> tuple:
    """Descrição, valor da primeira parcela, número dela e primeiro mês: o que a aba mostra."""
    primeira = parcelamento.primeira
    return (
        primeira.descricao,
        primeira.valor_centavos,
        str(primeira.parcela[0]),
        primeira.mes,
    )


def aba_parcelas_quebradas(parcelamentos: list[Parcelamento]) -> tuple[list[str], list[list]]:
    cabecalho = ["Descrição", "Valor", "Parcelas", "Primeiro mês", "Último mês", "Cartão", "Motivo"]
    linhas = [
        [
            p.descricao,
            reais(p.primeira.valor_centavos),
            p.faixa(),
            p.primeira.mes,
            p.ultima.mes,
            p.cartao,
            p.quebra,
        ]
        for p in parcelamentos
        if p.quebra
    ]
    return cabecalho, linhas


def aba_parcelamentos_sem_cartao(
    parcelamentos: list[Parcelamento], escolhidos: dict[tuple, str]
) -> tuple[list[str], list[list]]:
    """Os parcelamentos que nenhuma fatura trouxe: o dono escolhe o cartão, ou deixa em branco
    para ficarem na conta."""
    cabecalho = [
        "Descrição",
        "Valor",
        "Parcelas",
        "Primeiro mês",
        "Último mês",
        COLUNA_DO_CARTAO_ESCOLHIDO,
    ]
    linhas = []
    for p in parcelamentos:
        chave = chave_do_parcelamento(p)
        if p.cartao is None or chave in escolhidos:
            linhas.append(
                [
                    p.descricao,
                    reais(p.primeira.valor_centavos),
                    p.faixa(),
                    p.primeira.mes,
                    p.ultima.mes,
                    escolhidos.get(chave),
                ]
            )
    return cabecalho, linhas


def aba_ajustes_sem_cartao(escolhas: list[tuple]) -> tuple[list[str], list[list]]:
    """Os `Reajuste de fatura` que só o geral trouxe: o dono escolhe a fatura de que cartão o
    ajuste acerta. Cada escolha é o ajuste como o geral o trouxe e o cartão escolhido."""
    cabecalho = ["Mês", "Valor", COLUNA_DO_CARTAO_ESCOLHIDO]
    linhas = [[a.mes, reais(a.valor_centavos), cartao or None] for a, cartao in escolhas]
    return cabecalho, linhas


def aba_categoria_diferente(cruzamento: Cruzamento) -> tuple[list[str], list[list]]:
    """Os pares que casaram em tudo menos na categoria: vale a do geral."""
    cabecalho = [
        COLUNA_DA_PASTA,
        "Mês",
        "Data da compra",
        "Descrição",
        "Valor",
        "Categoria na fatura",
        "Categoria no geral (vale esta)",
    ]
    linhas = [
        [
            par.fatura.pasta,
            par.fatura.mes,
            data_br(par.fatura.data),
            par.fatura.descricao,
            reais(par.fatura.valor_centavos),
            par.fatura.categoria,
            par.geral.categoria,
        ]
        for par in cruzamento.pares
        if par.como == "sem categoria"
    ]
    return cabecalho, linhas


def _diferenca(nosso: int, do_mobills: int | None) -> float | None:
    return None if do_mobills is None else reais(nosso - do_mobills)


def aba_conferencia(meses: list[Mes], mobills: dict[str, int]) -> tuple[list[str], list[list]]:
    """Uma linha por mês, no mês da fatura. O dono preenche a coluna Mobills com o total do mês
    no Mobills (receitas - despesas) e roda de novo para ver a diferença."""
    cabecalho = [
        "Mês",
        "Receitas",
        "Despesas",
        "Ajustes de fatura",
        "Resultado",
        "Mobills",
        "Diferença",
        "Por quê",
    ]
    linhas = []
    for mes in meses:
        motivos = []
        for efeitos, o_que in (
            (mes.sem_par, ("linha de fatura sem par", "linhas de fatura sem par")),
            (mes.provaveis_abertos, ("par provável aberto", "pares prováveis abertos")),
            (
                mes.descartados,
                ("avulso pendente descartado", "avulsos pendentes descartados"),
            ),
        ):
            if efeitos:
                nome = o_que[0] if len(efeitos) == 1 else o_que[1]
                motivos.append(f"{len(efeitos)} {nome} ({reais_br(sum(efeitos))})")
        linhas.append(
            [
                mes.mes,
                reais(mes.receitas),
                reais(mes.despesas),
                reais(mes.ajustes),
                reais(mes.resultado),
                reais(mobills[mes.mes]) if mes.mes in mobills else None,
                _diferenca(mes.resultado, mobills.get(mes.mes)),
                "; ".join(motivos) or None,
            ]
        )
    return cabecalho, linhas


def aba_faturas(
    faturas: list[Fatura], mobills: dict[tuple[str, str], int]
) -> tuple[list[str], list[list]]:
    """O total de cada fatura, para comparar com a tela de faturas do Mobills."""
    cabecalho = ["Cartão", "Mês", "Compras", "Ajuste de fatura", "Total", "Mobills", "Diferença"]
    linhas = []
    for fatura in faturas:
        do_mobills = mobills.get((fatura.cartao, fatura.mes))
        linhas.append(
            [
                fatura.cartao,
                fatura.mes,
                reais(fatura.compras),
                reais(fatura.ajuste),
                reais(fatura.total),
                None if do_mobills is None else reais(do_mobills),
                _diferenca(fatura.total, do_mobills),
            ]
        )
    return cabecalho, linhas


def escrever_planilha(caminho: Path, abas: dict[str, tuple[list[str], list[list]]]) -> None:
    """Cada aba com o cabeçalho em negrito, larguras pelo conteúdo e valores em reais."""
    livro = Workbook()
    livro.remove(livro.active)
    for nome, (cabecalho, linhas) in abas.items():
        aba = livro.create_sheet(nome)
        aba.append(cabecalho)
        for celula in aba[1]:
            celula.font = Font(bold=True)
        for valores in linhas:
            aba.append(valores)
        aba.freeze_panes = "A2"
        for indice, titulo in enumerate(cabecalho, start=1):
            coluna = [titulo, *(v[indice - 1] for v in linhas)]
            largura = max(len(str(v)) for v in coluna if v is not None)
            aba.column_dimensions[get_column_letter(indice)].width = min(max(largura + 2, 8), 60)
            if any(isinstance(v[indice - 1], float) for v in linhas):
                for (celula,) in aba.iter_rows(min_row=2, min_col=indice, max_col=indice):
                    celula.number_format = FORMATO_DE_REAIS
    caminho.parent.mkdir(parents=True, exist_ok=True)
    livro.save(caminho)
