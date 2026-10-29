"""Leitura dos exports do Mobills (tickets 01 e 07): o geral, com a aba `Receitas e Despesas`, e as
faturas por cartão e mês, uma pasta por cartão com arquivos `<cartão>_-_<AAAAMM>.xlsx`."""

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ABA_DO_GERAL = "Receitas e Despesas"
ROTULO_DA_FATURA = re.compile(r"(\d{4})(\d{2})\.xlsx$", re.IGNORECASE)
PARCELA = re.compile(r"\s*\((\d+)/(\d+)\)\s*$")


class ArquivoInvalido(Exception):
    """Um arquivo do export, ou a planilha de revisão, que o importador não sabe ler."""


@dataclass(frozen=True)
class LinhaDoGeral:
    posicao: int  # a ordem no export inteiro
    arquivo: str
    linha: int  # a linha na aba, para o dono achar
    data: date
    descricao: str
    valor_centavos: int  # com o sinal do Mobills: despesa negativa
    conta: str
    situacao: str
    categoria: str
    tags: tuple[str, ...]
    parcela: tuple[int, int] | None

    @property
    def mes(self) -> str:
        return f"{self.data:%Y-%m}"


@dataclass(frozen=True)
class LinhaDeFatura:
    pasta: str
    mes: str  # o rótulo AAAAMM do arquivo, que é o mês da fatura
    arquivo: str
    data: date  # a data da compra
    descricao: str
    categoria: str
    valor_centavos: int  # como na fatura: compra positiva, crédito negativo


def ler_geral(pasta: Path) -> list[LinhaDoGeral]:
    """As linhas de lançamento de todos os XLSX da pasta, na ordem dos arquivos e das linhas. A
    linha em branco e a de total de cada aba não têm data e ficam de fora."""
    linhas: list[LinhaDoGeral] = []
    for arquivo in sorted(pasta.glob("*.xlsx")):
        livro = load_workbook(arquivo, read_only=True, data_only=True)
        if ABA_DO_GERAL not in livro.sheetnames:
            raise ArquivoInvalido(f"{arquivo.name}: falta a aba '{ABA_DO_GERAL}'")
        for numero, valores in _linhas_com_cabecalho(livro[ABA_DO_GERAL], arquivo):
            dia = data(valores.get("Data"))
            if dia is None:
                continue
            descricao, parcela = _sem_parcela(texto(valores.get("Descrição")))
            linhas.append(
                LinhaDoGeral(
                    posicao=len(linhas),
                    arquivo=arquivo.name,
                    linha=numero,
                    data=dia,
                    descricao=descricao,
                    valor_centavos=centavos(valores.get("Valor"), arquivo.name),
                    conta=texto(valores.get("Conta")),
                    situacao=texto(valores.get("Situação")),
                    categoria=texto(valores.get("Categoria")),
                    tags=_tags(valores.get("Tags")),
                    parcela=parcela,
                )
            )
        livro.close()
    return linhas


def ler_faturas(pasta: Path) -> list[LinhaDeFatura]:
    """As compras de todas as faturas, uma subpasta por cartão."""
    linhas: list[LinhaDeFatura] = []
    for subpasta in sorted(p for p in pasta.iterdir() if p.is_dir()):
        for arquivo in sorted(subpasta.glob("*.xlsx")):
            mes = _mes_do_rotulo(arquivo)
            livro = load_workbook(arquivo, read_only=True, data_only=True)
            for _, valores in _linhas_com_cabecalho(livro.worksheets[0], arquivo):
                dia = data(valores.get("Data"))
                if dia is None:
                    continue
                linhas.append(
                    LinhaDeFatura(
                        pasta=subpasta.name,
                        mes=mes,
                        arquivo=arquivo.name,
                        data=dia,
                        descricao=_sem_parcela(texto(valores.get("Descrição")))[0],
                        categoria=texto(valores.get("Categoria")),
                        valor_centavos=centavos(valores.get("Valor"), arquivo.name),
                    )
                )
            livro.close()
    return linhas


def pastas_de_fatura(pasta: Path) -> dict[str, list[str]]:
    """Cada pasta de cartão com os meses das faturas dela, inclusive as vazias."""
    pastas: dict[str, list[str]] = {}
    for subpasta in sorted(p for p in pasta.iterdir() if p.is_dir()):
        pastas[subpasta.name] = sorted(_mes_do_rotulo(a) for a in subpasta.glob("*.xlsx"))
    return pastas


def _mes_do_rotulo(arquivo: Path) -> str:
    """O mês da fatura, `AAAA-MM`, do rótulo `AAAAMM` no fim do nome do arquivo."""
    rotulo = ROTULO_DA_FATURA.search(arquivo.name)
    if rotulo is None:
        raise ArquivoInvalido(f"{arquivo.name}: falta o rótulo AAAAMM no fim do nome")
    return f"{rotulo[1]}-{rotulo[2]}"


def _linhas_com_cabecalho(aba, arquivo: Path):
    linhas = aba.iter_rows(values_only=True)
    cabecalho = [texto(c) for c in next(linhas, ())]
    if "Data" not in cabecalho or "Valor" not in cabecalho:
        raise ArquivoInvalido(f"{arquivo.name}: cabeçalho sem 'Data' e 'Valor'")
    for numero, valores in enumerate(linhas, start=2):
        yield numero, dict(zip(cabecalho, valores))


def texto(valor) -> str:
    """O conteúdo da célula como texto, sem espaços sobrando."""
    return "" if valor is None else " ".join(str(valor).split())


def data(valor) -> date | None:
    """A data da célula, em data ou em texto `dd/mm/aaaa`; `None` se não for data."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(texto(valor), "%d/%m/%Y").date()
    except ValueError:
        return None


def centavos(valor, onde: str) -> int:
    """O valor em centavos: número (`-120.35`) ou texto (`R$ 1.234,56`, `-1.234,56`, `12.5`)."""
    if isinstance(valor, int | float):
        return int((Decimal(str(valor)) * 100).quantize(Decimal(1)))
    numero = texto(valor).replace("R$", "").replace(" ", "")
    negativo = numero.startswith("-")
    numero = numero.lstrip("-")
    if "," in numero:
        numero = numero.replace(".", "").replace(",", ".")
    try:
        resultado = int((Decimal(numero) * 100).quantize(Decimal(1)))
    except ArithmeticError:
        raise ArquivoInvalido(f"{onde}: valor que não é número: {valor!r}") from None
    return -resultado if negativo else resultado


def _sem_parcela(descricao: str) -> tuple[str, tuple[int, int] | None]:
    """A descrição sem o sufixo `(n/total)` e a parcela que ele dizia."""
    achado = PARCELA.search(descricao)
    if achado is None:
        return descricao, None
    return descricao[: achado.start()], (int(achado[1]), int(achado[2]))


def _tags(valor) -> tuple[str, ...]:
    # O separador de várias tags não apareceu nos exports (ticket 01); a vírgula é a aposta, a
    # conferir no export real (ticket 89).
    return tuple(t.strip() for t in texto(valor).split(",") if t.strip())
